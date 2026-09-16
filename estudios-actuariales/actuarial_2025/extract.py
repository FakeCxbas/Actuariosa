from pathlib import Path
from pypdf import PdfReader
import openpyxl,json
p=Path(__file__).parent
for f in (p/'fuentes').iterdir():
 if f.suffix.lower()=='.pdf':
  r=PdfReader(f); (f.with_suffix('.txt')).write_text('\n'.join(f'\n--- PAGINA {i+1} ---\n'+(pg.extract_text(extraction_mode='layout') or '') for i,pg in enumerate(r.pages)),encoding='utf-8')
 if f.suffix.lower()=='.xlsx':
  w=openpyxl.load_workbook(f,data_only=True)
  data={s.title:[{'row':r[0].row,'cells':{c.column_letter:c.value for c in r if c.value is not None}} for r in s if any(c.value is not None for c in r)] for s in w}
  f.with_suffix('.json').write_text(json.dumps(data,ensure_ascii=False,default=str),encoding='utf-8')
  print(f.name,[(s.title,s.max_row,s.max_column) for s in w])
